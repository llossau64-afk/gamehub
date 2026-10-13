"""Local document parsing (PDF, DOCX, CSV, TXT, Markdown, source code)."""
from __future__ import annotations

import csv
import io
import logging
from pathlib import Path

log = logging.getLogger("velqin.documents")

MAX_TEXT_CHARS = 2_000_000


def extract_text(path: Path, type_key: str) -> str:
    try:
        if type_key == "pdf":
            return _pdf(path)
        if type_key == "docx":
            return _docx(path)
        if type_key == "csv":
            return _csv(path)
        return _plain(path)
    except Exception as e:  # noqa: BLE001 — parsing untrusted files must never crash the server
        log.warning("Text extraction failed for %s (%s): %s", path.name, type_key, type(e).__name__)
        raise ValueError("The file could not be read. It may be damaged or password protected.") from e


def _plain(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)[:MAX_TEXT_CHARS]
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")[:MAX_TEXT_CHARS]


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:  # noqa: BLE001
            raise ValueError("encrypted") from e
    parts: list[str] = []
    total = 0
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        parts.append(f"\n[Page {i + 1}]\n{text}")
        total += len(text)
        if total > MAX_TEXT_CHARS:
            break
    out = "".join(parts).strip()
    if not out.replace("[Page", "").strip():
        return "(This PDF contains no extractable text — it may be a scan. Upload pages as images for the vision model.)"
    return out[:MAX_TEXT_CHARS]


def _docx(path: Path) -> str:
    import docx  # python-docx

    d = docx.Document(str(path))
    parts: list[str] = []
    for para in d.paragraphs:
        if para.text.strip():
            style = (para.style.name or "").lower() if para.style is not None else ""
            prefix = "#" * int(style[-1]) + " " if style.startswith("heading") and style[-1:].isdigit() else ""
            parts.append(prefix + para.text)
    for t_index, table in enumerate(d.tables):
        parts.append(f"\n[Table {t_index + 1}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)[:MAX_TEXT_CHARS]


def _csv(path: Path) -> str:
    text = _plain(path)
    try:
        dialect = csv.Sniffer().sniff(text[:5000])
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = []
    for i, row in enumerate(reader):
        if i >= 2000:
            rows.append(f"... (truncated, file has more rows)")
            break
        rows.append(" | ".join(row))
    return "\n".join(rows)


def chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    text = text.strip()
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            # prefer to cut at a paragraph or sentence boundary
            cut = max(text.rfind("\n\n", start + size // 2, end), text.rfind(". ", start + size // 2, end))
            if cut > start:
                end = cut + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]
